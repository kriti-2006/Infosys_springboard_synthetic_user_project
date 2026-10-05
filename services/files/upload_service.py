"""
Upload pipeline:  validate (ext + MIME + size + magic bytes + structure) -> store (hashed name, outside web root)
                  -> extract text -> normalise -> persist on the experiment (used by services/experiment_context.py).
Raw files are NEVER sent to the LLM; only bounded, normalised text is. Allowed: pdf, docx, txt, csv.
"""
import csv
import hashlib
import io
import re
import zipfile
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from config import settings
from db.models import Experiment, ExperimentDocument
from services.errors import Conflict, InvalidInput, NotFound
from utils.logging import get_logger, log_event

log = get_logger("uploads")
ALLOWED = {
    "pdf": {"application/pdf"},
    "docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/octet-stream"},
    "txt": {"text/plain"},
    "csv": {"text/csv", "application/csv", "text/plain", "application/vnd.ms-excel"},
}
MAX_TEXT_CHARS = 200_000
MAX_PDF_PAGES = 200
MAX_CSV_ROWS = 500
MAX_ZIP_UNCOMPRESSED = 50 * 1024 * 1024


def _safe_name(name: str) -> str:
    name = Path(name or "upload").name
    return re.sub(r"[^A-Za-z0-9._\- ]", "_", name)[:200] or "upload"


def validate(filename: str, content_type: str | None, data: bytes) -> str:
    ext = Path(filename or "").suffix.lower().lstrip(".")
    if ext not in ALLOWED:
        raise InvalidInput(f"file type .{ext or '?'} not allowed (allowed: {', '.join(sorted(ALLOWED))})")
    ctype = (content_type or "").split(";")[0].strip().lower()
    if ctype and ctype not in ALLOWED[ext]:
        raise InvalidInput(f"MIME type {ctype} does not match .{ext}")
    if len(data) == 0:
        raise InvalidInput("file is empty")
    if len(data) > settings.MAX_UPLOAD_MB * 1024 * 1024:
        raise InvalidInput(f"file exceeds {settings.MAX_UPLOAD_MB} MB limit")
    if ext == "pdf" and not data.startswith(b"%PDF-"):
        raise InvalidInput("file content is not a PDF")
    if ext == "docx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = z.namelist()
                if "[Content_Types].xml" not in names or not any(n.startswith("word/") for n in names):
                    raise InvalidInput("file content is not a Word document")
                if sum(i.file_size for i in z.infolist()) > MAX_ZIP_UNCOMPRESSED:
                    raise InvalidInput("document expands too large (possible zip bomb)")
        except zipfile.BadZipFile:
            raise InvalidInput("file content is not a Word document")
    if ext in ("txt", "csv"):
        if b"\x00" in data[:4096]:
            raise InvalidInput("binary content in a text file")
        try:
            data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise InvalidInput("text files must be UTF-8")
    return ext


def extract_text(ext: str, data: bytes) -> str:
    if ext == "pdf":
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((pg.extract_text() or "") for pg in reader.pages[:MAX_PDF_PAGES])
        except Exception as e:  # noqa: BLE001
            raise InvalidInput(f"could not read PDF: {type(e).__name__}")
    elif ext == "docx":
        import docx
        try:
            d = docx.Document(io.BytesIO(data))
            text = "\n".join(p.text for p in d.paragraphs)
            for t in d.tables:
                text += "\n" + "\n".join(" | ".join(c.text for c in r.cells) for r in t.rows)
        except Exception as e:  # noqa: BLE001
            raise InvalidInput(f"could not read DOCX: {type(e).__name__}")
    elif ext == "csv":
        rows = csv.reader(io.StringIO(data.decode("utf-8-sig")))
        text = "\n".join(" | ".join(r) for _, r in zip(range(MAX_CSV_ROWS), rows))
    else:
        text = data.decode("utf-8-sig")
    text = re.sub(r"[ \t]+", " ", text.replace("\r", ""))
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text[:MAX_TEXT_CHARS]


async def save_document(s: AsyncSession, experiment_id: str, filename: str, content_type: str | None, data: bytes) -> ExperimentDocument:
    if not await s.get(Experiment, experiment_id):
        raise NotFound("experiment not found")
    ext = validate(filename, content_type, data)
    sha = hashlib.sha256(data).hexdigest()
    if await s.scalar(select(ExperimentDocument.id).where(ExperimentDocument.experiment_id == experiment_id, ExperimentDocument.sha256 == sha)):
        raise Conflict("this file was already uploaded to the experiment")
    text = extract_text(ext, data)
    folder = Path(settings.UPLOAD_DIR) / experiment_id
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{sha}.{ext}"  # server-generated name: user filename never touches the filesystem path
    path.write_bytes(data)
    doc = ExperimentDocument(experiment_id=experiment_id, filename=_safe_name(filename), mime_type=(content_type or ext)[:100],
                             size_bytes=len(data), sha256=sha, storage_path=str(path), extracted_text=text)
    s.add(doc)
    try:
        await s.commit()
    except IntegrityError:
        await s.rollback()
        raise Conflict("this file was already uploaded to the experiment")
    log_event(log, "document_uploaded", experiment_id=experiment_id, document_id=doc.id, ext=ext, chars=len(text))
    return doc
